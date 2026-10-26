from openpyxl.styles import PatternFill

def get_failure_rate_fill(percentage):
    """
    Returns an openpyxl PatternFill based on the failure rate percentage.
    Green: <= 10
    Orange: > 10 and <= 33
    Red: > 33
    """
    if percentage > 33:
        return PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")  # Red
    elif percentage > 10 and percentage <= 33:
        return PatternFill(start_color="FFA500", end_color="FFA500", fill_type="solid")  # Orange
    else:
        return PatternFill(start_color="00FF00", end_color="00FF00", fill_type="solid")  # Green
